"""Extract plain text from a source document.

Supports Markdown/plain text, PDF, and Word (.docx). Extraction is deliberately
no-frills: we want predictable text to chunk, not layout fidelity. If layout or
tables matter later, `unstructured` is the heavier upgrade path.
"""

from __future__ import annotations

from pathlib import Path

# File extensions we know how to read, grouped by extractor.
TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".rst"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | PDF_SUFFIXES | DOCX_SUFFIXES


class UnsupportedDocument(ValueError):
    """Raised when a file's type has no registered extractor."""


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_SUFFIXES


def extract_text(path: Path) -> str:
    """Return the plain text of `path`, dispatching on file extension.

    Raises UnsupportedDocument for unknown types so callers can skip/report
    rather than silently ingesting garbage.
    """
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _extract_text_file(path)
    if suffix in PDF_SUFFIXES:
        return _extract_pdf(path)
    if suffix in DOCX_SUFFIXES:
        return _extract_docx(path)
    raise UnsupportedDocument(f"no extractor for {path.suffix!r} ({path.name})")


def _extract_text_file(path: Path) -> str:
    # utf-8 with a lenient fallback; docs occasionally carry stray bytes.
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)
